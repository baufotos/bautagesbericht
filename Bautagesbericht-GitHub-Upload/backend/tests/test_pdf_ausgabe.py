"""PDF-Ausgabe auf beiden Betriebsarten — Word im Büro, LibreOffice im Netz.

Warum das eine eigene Testreihe hat: Die App läuft an zwei Orten mit
verschiedener Ausstattung. Auf dem Bürorechner erzeugt Microsoft Word das PDF,
im Linux-Container von Render gibt es kein Word — dort muss LibreOffice
einspringen. Beide Wege sind von hier aus nicht gleichzeitig prüfbar: Auf
diesem Windows-Rechner gibt es kein LibreOffice, im Container kein Word.

Deshalb werden die Weichen mit untergeschobenen Doppeln geprüft
(``_ist_windows``, ``shutil.which``, ``subprocess.run``) und nur die
tatsächlich vorhandene Umwandlung zusätzlich echt durchgeführt. Das ist die
einzige Art, die Linux-Weiche zu beweisen, ohne Linux zu haben.

Die wichtigste Zusicherung steckt in Abschnitt 2: Auf Windows wird ZUERST
Word gefragt. Das Windows-Paket ist im Einsatz, und an seinem Verhalten darf
sich durch den neuen Server-Weg nichts ändern.
"""
import io
import os
import re
import subprocess
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BACKEND = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND))

from app.services import word_pdf  # noqa: E402

ok = 0
fehler: list[str] = []


def pruefe(bedingung, text):
    global ok
    if bedingung:
        ok += 1
    else:
        fehler.append(text)


def gleich(ist, soll, text):
    pruefe(ist == soll, f"{text}: erwartet {soll!r}, war {ist!r}")


# ── Doppel für die Umgebung ────────────────────────────────────────────────
#
# Gemerkt und am Ende zurückgesetzt, damit die echte Umwandlung in Abschnitt 6
# nicht auf einem untergeschobenen Doppel landet.
ECHT_RUN = subprocess.run
ECHT_WHICH = word_pdf.shutil.which
ECHT_WINDOWS = word_pdf._ist_windows

PDF_PROBE = b"%PDF-1.7\nprobe\n"


#: Findet den Zielpfad auch dort, wo er nicht als eigenes Befehlswort steht.
#: Nötig, weil die beiden Wege ihn verschieden übergeben: LibreOffice bekommt
#: die Quelldatei als Wort auf der Befehlszeile, Word dagegen ein PowerShell-
#: Skript, in dem Quelle und Ziel als Zuweisungen im Text stehen.
_ZIELPFAD = re.compile(r"[^\"'\s]+bericht\.pdf")


def _lege_pdf_ab(befehl):
    """Ahmt nach, was beide Programme tun: bericht.pdf neben die Quelle legen."""
    for teil in befehl:
        text = str(teil)
        p = Path(text)
        if p.suffix == ".docx" and p.parent.is_dir():
            (p.parent / "bericht.pdf").write_bytes(PDF_PROBE)
            return
        treffer = _ZIELPFAD.search(text)
        if treffer:
            ziel = Path(treffer.group(0))
            if ziel.parent.is_dir():
                ziel.write_bytes(PDF_PROBE)
                return


def _ergebnis(returncode=0, stdout="", stderr=""):
    e = type("Ergebnis", (), {})()
    e.returncode = returncode
    e.stdout = stdout
    e.stderr = stderr
    return e


class Lauf:
    """Merkt sich, womit ``subprocess.run`` aufgerufen wurde."""

    def __init__(self, erfolg=True, schreibt_pdf=True):
        self.erfolg = erfolg
        self.schreibt_pdf = schreibt_pdf
        self.befehle = []

    def __call__(self, befehl, **kw):
        self.befehle.append(list(befehl))
        if self.schreibt_pdf:
            _lege_pdf_ab(befehl)
        return _ergebnis(0 if self.erfolg else 1,
                         "" if self.erfolg else "Fehlermeldung des Programms")


class WordScheitertLibreOfficeNicht:
    """Word meldet Fehler, LibreOffice danach Erfolg — der Windows-Rückfall."""

    def __init__(self):
        self.befehle = []

    def __call__(self, befehl, **kw):
        self.befehle.append(list(befehl))
        if "powershell" in str(befehl[0]).lower():
            return _ergebnis(1, "Word nicht installiert")
        _lege_pdf_ab(befehl)
        return _ergebnis(0)


def stelle(windows, libreoffice):
    word_pdf._ist_windows = lambda: windows
    word_pdf.shutil.which = lambda name: (
        libreoffice if name in word_pdf._LIBREOFFICE_NAMEN else None
    )


def zuruecksetzen():
    subprocess.run = ECHT_RUN
    word_pdf.shutil.which = ECHT_WHICH
    word_pdf._ist_windows = ECHT_WINDOWS


DOCX = b"PK\x03\x04 nur ein Platzhalter"

print("1. Auskunft: was kann dieser Rechner?")
stelle(windows=False, libreoffice=None)
gleich(word_pdf.word_vorhanden(), False, "kein Word auf Linux")
gleich(word_pdf.libreoffice_pfad(), None, "kein LibreOffice gefunden")
gleich(word_pdf.pdf_moeglich(), False, "ohne beides kein PDF")
gleich(word_pdf.pdf_weg(), "keins", "Weg heisst keins")

stelle(windows=False, libreoffice="/usr/bin/soffice")
gleich(word_pdf.libreoffice_pfad(), "/usr/bin/soffice", "LibreOffice gefunden")
gleich(word_pdf.pdf_moeglich(), True, "mit LibreOffice geht PDF")
gleich(word_pdf.pdf_weg(), "libreoffice", "Weg heisst libreoffice")

print("2. Windows nimmt WEITERHIN zuerst Word (unveraenderte Buero-Logik)")
stelle(windows=True, libreoffice="C:/LO/soffice.exe")
lauf = Lauf()
subprocess.run = lauf
word_pdf.nach_pdf(DOCX)
gleich(len(lauf.befehle), 1, "genau ein Aufruf")
pruefe("powershell" in lauf.befehle[0][0].lower(),
       f"Word ueber PowerShell, war {lauf.befehle[0][0]!r}")
pruefe(not any("soffice" in str(t) for t in lauf.befehle[0]),
       "LibreOffice wurde NICHT angefasst, obwohl es da ist")

print("3. Linux nimmt LibreOffice — mit den richtigen Schaltern")
stelle(windows=False, libreoffice="/usr/bin/soffice")
lauf = Lauf()
subprocess.run = lauf
ergebnis = word_pdf.nach_pdf(DOCX)
gleich(ergebnis[:5], b"%PDF-", "PDF-Bytes kommen zurueck")
befehl = lauf.befehle[0]
gleich(befehl[0], "/usr/bin/soffice", "richtige Programmdatei")
for schalter in ("--headless", "--invisible", "--norestore", "--nolockcheck",
                 "--nodefault", "--nofirststartwizard", "--convert-to",
                 "--outdir"):
    pruefe(schalter in befehl, f"Schalter {schalter} fehlt")
gleich(befehl[befehl.index("--convert-to") + 1], "pdf:writer_pdf_Export",
       "Writer-Filter ausdruecklich benannt")
profil = [t for t in befehl if t.startswith("-env:UserInstallation=")]
gleich(len(profil), 1, "eigenes Benutzerprofil gesetzt")
pruefe(profil and profil[0].startswith("-env:UserInstallation=file:///"),
       f"Profil als file:-URL, war {profil}")
pruefe(str(word_pdf._profil_ordner()).endswith(str(os.getpid())),
       "Profilname enthaelt die Prozessnummer")

print("4. Linux ohne LibreOffice: klare Absage statt kaputter Datei")
stelle(windows=False, libreoffice=None)
try:
    word_pdf.nach_pdf(DOCX)
    fehler.append("haette werfen muessen")
except word_pdf.PdfNichtMoeglich as f:
    text = str(f)
    pruefe("Word-Dokument" in text, f"nennt die Ausweichloesung: {text}")
    pruefe("herunterladen" in text, f"sagt was zu tun ist: {text}")

print("5. Windows ohne Word: frueher Fehler, jetzt LibreOffice — falls da")
stelle(windows=True, libreoffice=None)
subprocess.run = Lauf(erfolg=False, schreibt_pdf=False)
try:
    word_pdf.nach_pdf(DOCX)
    fehler.append("ohne Word und ohne LibreOffice haette es werfen muessen")
except word_pdf.PdfNichtMoeglich as f:
    pruefe("Word" in str(f), f"alte Windows-Meldung bleibt: {f}")

stelle(windows=True, libreoffice="C:/LO/soffice.exe")
lauf = WordScheitertLibreOfficeNicht()
subprocess.run = lauf
ergebnis = word_pdf.nach_pdf(DOCX)
gleich(ergebnis[:5], b"%PDF-", "LibreOffice rettet den Windows-Fall")
gleich(len(lauf.befehle), 2, "erst Word, dann LibreOffice")
pruefe("powershell" in lauf.befehle[0][0].lower(), "Word war zuerst dran")
pruefe("soffice" in lauf.befehle[1][0].lower(), "LibreOffice war zweiter")

zuruecksetzen()

print("6. Echte Umwandlung mit dem, was auf DIESEM Rechner da ist")
if not word_pdf.pdf_moeglich():
    print("   Weder Word noch LibreOffice erreichbar — uebersprungen.")
else:
    from docx import Document

    dok = Document()
    dok.add_paragraph("Bautagesbericht — Probe")
    dok.add_paragraph("Zweite Zeile, damit das Blatt nicht leer ist.")
    puffer = io.BytesIO()
    dok.save(puffer)
    pdf = word_pdf.nach_pdf(puffer.getvalue())
    gleich(pdf[:5], b"%PDF-", f"echtes PDF ueber {word_pdf.pdf_weg()}")
    pruefe(len(pdf) > 800, f"PDF hat Inhalt ({len(pdf)} Bytes)")
    print(f"   Weg: {word_pdf.pdf_weg()}, {len(pdf)} Bytes")

zuruecksetzen()

print()
print(f"{ok} Pruefungen ok, {len(fehler)} Fehler")
if fehler:
    print("FEHLER:")
    for f in fehler:
        print(" -", f)
    sys.exit(1)
